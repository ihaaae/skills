#!/usr/bin/env -S uv run --script
# /// script
# requires-python = ">=3.9"
# dependencies = ["openpyxl>=3.1.5"]
# ///

from __future__ import annotations

import argparse
import csv
import json
import math
import sys
from collections import OrderedDict
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


EXCEL_EXTENSIONS = {".xlsx", ".xlsm", ".xltx", ".xltm"}
CSV_EXTENSIONS = {".csv", ".tsv"}
BLANK_STRINGS = {"", "null", "none", "nan"}


class CompareError(Exception):
    pass


@dataclass
class Table:
    headers: List[str]
    rows: List[OrderedDict]
    sheet: str


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Compare two table files and write a deterministic diff workbook."
    )
    parser.add_argument("old_file", help="Old/baseline .xlsx, .csv, or .tsv file")
    parser.add_argument("new_file", help="New/comparison .xlsx, .csv, or .tsv file")
    parser.add_argument(
        "--mode",
        choices=["key", "position"],
        default="key",
        help="Compare by primary key or by row/column position. Default: key",
    )
    parser.add_argument("--sheet", help="Excel worksheet name. Defaults to active sheet.")
    parser.add_argument(
        "--header-row",
        type=int,
        default=1,
        help="1-based header row for Excel files. Default: 1",
    )
    parser.add_argument("--key", help="Primary key column for --mode key")
    parser.add_argument(
        "--ignore-cols",
        default="",
        help="Comma-separated column names to ignore, e.g. 更新时间,导出时间",
    )
    parser.add_argument(
        "--tolerance",
        type=float,
        default=0.0,
        help="Numeric tolerance for treating numbers as equal. Default: 0",
    )
    parser.add_argument(
        "--out",
        default="diff.xlsx",
        help="Output Excel diff workbook. Default: diff.xlsx",
    )
    parser.add_argument(
        "--json",
        dest="json_out",
        help="Optional JSON summary output path.",
    )
    return parser.parse_args()


def split_csv(value: str) -> List[str]:
    return [part.strip() for part in value.split(",") if part.strip()]


def as_header(value: Any, index: int) -> str:
    if value is None or str(value).strip() == "":
        return f"Column{index}"
    return str(value).strip()


def dedupe_headers(headers: Sequence[str]) -> List[str]:
    counts: Dict[str, int] = {}
    output: List[str] = []
    for header in headers:
        counts[header] = counts.get(header, 0) + 1
        if counts[header] == 1:
            output.append(header)
        else:
            output.append(f"{header}_{counts[header]}")
    return output


def read_table(path: Path, sheet: Optional[str], header_row: int) -> Table:
    suffix = path.suffix.lower()
    if suffix in EXCEL_EXTENSIONS:
        return read_excel(path, sheet, header_row)
    if suffix in CSV_EXTENSIONS:
        return read_delimited(path)
    raise CompareError(f"Unsupported file extension for {path}. Use .xlsx, .xlsm, .csv, or .tsv.")


def read_excel(path: Path, sheet: Optional[str], header_row: int) -> Table:
    if header_row < 1:
        raise CompareError("--header-row must be >= 1")
    try:
        workbook = load_workbook(path, data_only=False, read_only=False)
    except Exception as exc:
        raise CompareError(f"Could not read Excel file {path}: {exc}") from exc

    if sheet:
        if sheet not in workbook.sheetnames:
            raise CompareError(f"Sheet {sheet!r} not found in {path}. Available: {', '.join(workbook.sheetnames)}")
        worksheet = workbook[sheet]
    else:
        worksheet = workbook.active

    raw_headers = [cell.value for cell in worksheet[header_row]]
    headers = dedupe_headers([as_header(value, i + 1) for i, value in enumerate(raw_headers)])
    rows: List[OrderedDict] = []
    for excel_row in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
        if not any(value is not None and str(value).strip() != "" for value in excel_row):
            continue
        row = OrderedDict()
        for index, header in enumerate(headers):
            row[header] = excel_row[index] if index < len(excel_row) else None
        rows.append(row)
    return Table(headers=headers, rows=rows, sheet=worksheet.title)


def read_delimited(path: Path) -> Table:
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    try:
        with path.open("r", newline="", encoding="utf-8-sig") as handle:
            reader = csv.DictReader(handle, delimiter=delimiter)
            if not reader.fieldnames:
                raise CompareError(f"No header row found in {path}")
            headers = dedupe_headers([as_header(value, i + 1) for i, value in enumerate(reader.fieldnames)])
            rows = []
            for raw in reader:
                row = OrderedDict()
                for header in headers:
                    row[header] = raw.get(header)
                if any(value is not None and str(value).strip() != "" for value in row.values()):
                    rows.append(row)
    except UnicodeDecodeError as exc:
        raise CompareError(f"Could not decode {path} as UTF-8 CSV/TSV: {exc}") from exc
    return Table(headers=headers, rows=rows, sheet=path.name)


def normalize(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, str):
        stripped = value.strip()
        if stripped.lower() in BLANK_STRINGS:
            return ""
        return stripped
    return value


def to_float(value: Any) -> Optional[float]:
    normalized = normalize(value)
    if normalized == "":
        return None
    if isinstance(normalized, bool):
        return None
    try:
        return float(normalized)
    except (TypeError, ValueError):
        return None


def values_equal(old: Any, new: Any, tolerance: float) -> bool:
    old_norm = normalize(old)
    new_norm = normalize(new)
    if old_norm == "" and new_norm == "":
        return True
    old_float = to_float(old_norm)
    new_float = to_float(new_norm)
    if old_float is not None and new_float is not None:
        return math.isclose(old_float, new_float, abs_tol=tolerance, rel_tol=0.0)
    return old_norm == new_norm


def index_by_key(table: Table, key: str) -> Dict[str, OrderedDict]:
    if key not in table.headers:
        raise CompareError(f"Key column {key!r} not found. Available columns: {', '.join(table.headers)}")
    indexed: Dict[str, OrderedDict] = {}
    duplicates: List[str] = []
    for row_number, row in enumerate(table.rows, start=2):
        key_value = str(normalize(row.get(key)))
        if key_value == "":
            raise CompareError(f"Blank key value in column {key!r} near data row {row_number}")
        if key_value in indexed:
            duplicates.append(key_value)
        indexed[key_value] = row
    if duplicates:
        preview = ", ".join(duplicates[:10])
        raise CompareError(f"Duplicate key values in column {key!r}: {preview}")
    return indexed


def compare_by_key(old: Table, new: Table, key: str, ignore_cols: Sequence[str], tolerance: float) -> Dict[str, Any]:
    old_index = index_by_key(old, key)
    new_index = index_by_key(new, key)
    ignored = set(ignore_cols)
    ignored.add(key)
    comparable_columns = [
        column for column in old.headers if column in new.headers and column not in ignored
    ]

    added_keys = sorted(set(new_index) - set(old_index))
    deleted_keys = sorted(set(old_index) - set(new_index))
    common_keys = sorted(set(old_index) & set(new_index))

    changed: List[Dict[str, Any]] = []
    changed_keys = set()
    for key_value in common_keys:
        old_row = old_index[key_value]
        new_row = new_index[key_value]
        for column in comparable_columns:
            old_value = old_row.get(column)
            new_value = new_row.get(column)
            if values_equal(old_value, new_value, tolerance):
                continue
            old_float = to_float(old_value)
            new_float = to_float(new_value)
            delta = new_float - old_float if old_float is not None and new_float is not None else None
            changed.append(
                {
                    "key": key_value,
                    "field": column,
                    "old_value": old_value,
                    "new_value": new_value,
                    "delta": delta,
                }
            )
            changed_keys.add(key_value)

    unchanged = len(common_keys) - len(changed_keys)
    return {
        "mode": "key",
        "key": key,
        "added": [new_index[value] for value in added_keys],
        "deleted": [old_index[value] for value in deleted_keys],
        "changed": changed,
        "changed_keys": sorted(changed_keys),
        "summary": {
            "old_rows": len(old.rows),
            "new_rows": len(new.rows),
            "added": len(added_keys),
            "deleted": len(deleted_keys),
            "changed_rows": len(changed_keys),
            "changed_fields": len(changed),
            "unchanged": unchanged,
        },
    }


def compare_by_position(old: Table, new: Table, ignore_cols: Sequence[str], tolerance: float) -> Dict[str, Any]:
    ignored = set(ignore_cols)
    max_rows = max(len(old.rows), len(new.rows))
    max_cols = max(len(old.headers), len(new.headers))
    changed: List[Dict[str, Any]] = []
    changed_rows = set()

    for row_index in range(max_rows):
        old_row = old.rows[row_index] if row_index < len(old.rows) else OrderedDict()
        new_row = new.rows[row_index] if row_index < len(new.rows) else OrderedDict()
        for col_index in range(max_cols):
            old_header = old.headers[col_index] if col_index < len(old.headers) else f"Column{col_index + 1}"
            new_header = new.headers[col_index] if col_index < len(new.headers) else f"Column{col_index + 1}"
            field = old_header if old_header == new_header else f"{old_header} / {new_header}"
            if old_header in ignored or new_header in ignored:
                continue
            old_value = old_row.get(old_header)
            new_value = new_row.get(new_header)
            if values_equal(old_value, new_value, tolerance):
                continue
            changed.append(
                {
                    "row": row_index + 2,
                    "column": col_index + 1,
                    "field": field,
                    "old_value": old_value,
                    "new_value": new_value,
                }
            )
            changed_rows.add(row_index)

    return {
        "mode": "position",
        "added": [],
        "deleted": [],
        "changed": changed,
        "changed_keys": [],
        "summary": {
            "old_rows": len(old.rows),
            "new_rows": len(new.rows),
            "added": max(0, len(new.rows) - len(old.rows)),
            "deleted": max(0, len(old.rows) - len(new.rows)),
            "changed_rows": len(changed_rows),
            "changed_fields": len(changed),
            "unchanged": max_rows - len(changed_rows),
        },
    }


def stringify(value: Any) -> Any:
    if value is None:
        return ""
    return value


def append_table(sheet, headers: Sequence[str], rows: Iterable[Sequence[Any]]) -> None:
    sheet.append(list(headers))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for row in rows:
        sheet.append([stringify(value) for value in row])
    for column_cells in sheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        sheet.column_dimensions[column_cells[0].column_letter].width = min(max(max_length + 2, 12), 60)


def write_workbook(path: Path, result: Dict[str, Any], old: Table, new: Table) -> None:
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    summary_rows = [(key, value) for key, value in result["summary"].items()]
    append_table(summary_sheet, ["item", "count"], summary_rows)

    added_sheet = workbook.create_sheet("Added")
    added_headers = list(new.headers) + ["diff_type"]
    append_table(
        added_sheet,
        added_headers,
        ([row.get(header) for header in new.headers] + ["Added"] for row in result["added"]),
    )

    deleted_sheet = workbook.create_sheet("Deleted")
    deleted_headers = list(old.headers) + ["diff_type"]
    append_table(
        deleted_sheet,
        deleted_headers,
        ([row.get(header) for header in old.headers] + ["Deleted"] for row in result["deleted"]),
    )

    changed_sheet = workbook.create_sheet("Changed")
    if result["mode"] == "key":
        headers = ["key", "field", "old_value", "new_value", "delta", "diff_type"]
        rows = (
            [change["key"], change["field"], change["old_value"], change["new_value"], change.get("delta"), "Modified"]
            for change in result["changed"]
        )
    else:
        headers = ["row", "column", "field", "old_value", "new_value", "diff_type"]
        rows = (
            [change["row"], change["column"], change["field"], change["old_value"], change["new_value"], "Modified"]
            for change in result["changed"]
        )
    append_table(changed_sheet, headers, rows)
    workbook.save(path)


def json_safe(value: Any) -> Any:
    if isinstance(value, OrderedDict):
        return {key: json_safe(item) for key, item in value.items()}
    if isinstance(value, dict):
        return {key: json_safe(item) for key, item in value.items()}
    if isinstance(value, list):
        return [json_safe(item) for item in value]
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return value


def print_summary(result: Dict[str, Any], old_file: Path, new_file: Path, old: Table, new: Table, out: Path, json_out: Optional[Path]) -> None:
    summary = result["summary"]
    print("Comparing:")
    print(f"  Old file: {old_file}")
    print(f"  New file: {new_file}")
    print(f"  Old sheet/table: {old.sheet}")
    print(f"  New sheet/table: {new.sheet}")
    print(f"  Mode: {result['mode']}")
    if result.get("key"):
        print(f"  Key column: {result['key']}")
    print()
    print("Result:")
    print(f"  Rows in old file: {summary['old_rows']}")
    print(f"  Rows in new file: {summary['new_rows']}")
    print(f"  Added rows:      {summary['added']}")
    print(f"  Deleted rows:    {summary['deleted']}")
    print(f"  Changed rows:    {summary['changed_rows']}")
    print(f"  Changed fields:  {summary['changed_fields']}")
    print(f"  Unchanged rows:  {summary['unchanged']}")
    print()
    print("Output written to:")
    print(f"  {out}")
    if json_out:
        print(f"  {json_out}")


def main() -> int:
    args = parse_args()
    try:
        old_file = Path(args.old_file).expanduser()
        new_file = Path(args.new_file).expanduser()
        if not old_file.exists():
            raise CompareError(f"Old file not found: {old_file}")
        if not new_file.exists():
            raise CompareError(f"New file not found: {new_file}")
        if args.mode == "key" and not args.key:
            raise CompareError("--key is required when --mode key")

        ignore_cols = split_csv(args.ignore_cols)
        old = read_table(old_file, args.sheet, args.header_row)
        new = read_table(new_file, args.sheet, args.header_row)

        if args.mode == "key":
            result = compare_by_key(old, new, args.key, ignore_cols, args.tolerance)
        else:
            result = compare_by_position(old, new, ignore_cols, args.tolerance)

        out = Path(args.out).expanduser()
        write_workbook(out, result, old, new)
        json_out = Path(args.json_out).expanduser() if args.json_out else None
        if json_out:
            payload = {
                "files": {"old": str(old_file), "new": str(new_file)},
                "sheet": {"old": old.sheet, "new": new.sheet},
                "mode": result["mode"],
                "key": result.get("key"),
                "summary": result["summary"],
                "changes": result["changed"],
            }
            json_out.write_text(json.dumps(json_safe(payload), ensure_ascii=False, indent=2), encoding="utf-8")

        print_summary(result, old_file, new_file, old, new, out, json_out)
        has_differences = any(
            result["summary"][name] for name in ("added", "deleted", "changed_fields")
        )
        return 1 if has_differences else 0
    except CompareError as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
